from docx import Document

from .pdf_parser import extract_text as extract_pdf_text


def _extract_docx(file) -> str:
    doc = Document(file)
    lines = [p.text for p in doc.paragraphs]
    # Many resume templates put content in tables (for example a skills grid).
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                lines.extend(p.text for p in cell.paragraphs)
    return "\n".join(line for line in lines if line.strip()).strip()


def extract_document_text(file, two_column: bool = False) -> str:
    """Read text from a PDF or DOCX, chosen by file extension."""
    if file.name.lower().endswith(".docx"):
        return _extract_docx(file)
    return extract_pdf_text(file, two_column=two_column)