import io
import json
import urllib.error
from unittest.mock import MagicMock, patch

from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from analyser.services.section_parser import split_sections

from analyser.services.suggestions import SuggestionError, create_suggestions
from analyser.services.scorer import score_resume

import io
from docx import Document as WordDocument
from analyser.services.document_parser import extract_document_text


FAKE_KEY = "test-key-not-real"
SUGGEST_URL = "/api/suggestions/"  # check analyser/urls.py

LINES = [
    "Built a student attendance tracker using Python and SQLite",
    "Managed data for 120 students in the system",
    "Aspiring Python developer seeking a software role",
]
SECTIONS = {"projects": "\n".join(LINES[:2]), "summary": LINES[2]}
SOURCE = "\n".join(LINES)


def good_item(**overrides):
    item = {
        "original": LINES[0],
        "suggested_rewrite": "Developed a student attendance tracker in Python with an SQLite database",
        "reason": "Stronger action verb",
        "question": "How many students used it?",
    }
    item.update(overrides)
    return item


def api_response(text):
    """A fake urlopen() result containing Claude's text."""
    body = json.dumps({"content": [{"type": "text", "text": text}]}).encode()
    resp = MagicMock()
    resp.read.return_value = body
    resp.__enter__.return_value = resp
    return resp


def suggestions_response(*items):
    return api_response(json.dumps({"suggestions": list(items)}))


def http_error(code, message):
    body = json.dumps({"error": {"message": message}}).encode()
    return urllib.error.HTTPError("https://example.invalid", code, "error", {}, io.BytesIO(body))


@override_settings(CLAUDE_SUGGESTIONS_ENABLED=True, ANTHROPIC_API_KEY=FAKE_KEY, ANTHROPIC_MODEL="test-model")
@patch("analyser.services.suggestions.urllib.request.urlopen")
class SuggestionServiceTests(TestCase):
    def run_service(self):
        return create_suggestions(SECTIONS, ["docker"], ["git"])

    def test_valid_response_passes(self, urlopen):
        urlopen.return_value = suggestions_response(good_item())
        result = self.run_service()
        self.assertEqual(len(result["suggestions"]), 1)
        self.assertIn("draft_notice", result)

    def test_number_from_original_is_allowed(self, urlopen):
        urlopen.return_value = suggestions_response(good_item(
            original=LINES[1], suggested_rewrite="Managed records for 120 students"))
        self.assertEqual(len(self.run_service()["suggestions"]), 1)

    def test_invented_number_is_rejected(self, urlopen):
        urlopen.return_value = suggestions_response(good_item(
            suggested_rewrite="Built a tracker used by 500 students"))
        with self.assertRaisesMessage(SuggestionError, "unsupported details"):
            self.run_service()

    def test_invented_skill_is_rejected(self, urlopen):
        urlopen.return_value = suggestions_response(good_item(
            suggested_rewrite="Built a student attendance tracker deployed with Docker"))
        with self.assertRaisesMessage(SuggestionError, "unsupported details"):
            self.run_service()

    def test_unknown_source_line_is_rejected(self, urlopen):
        urlopen.return_value = suggestions_response(good_item(
            original="Led a team of engineers at a startup"))
        with self.assertRaisesMessage(SuggestionError, "did not match"):
            self.run_service()

    def test_extra_field_is_rejected(self, urlopen):
        urlopen.return_value = suggestions_response(good_item(confidence="high"))
        with self.assertRaisesMessage(SuggestionError, "unexpected response format"):
            self.run_service()

    def test_more_than_five_suggestions_rejected(self, urlopen):
        urlopen.return_value = suggestions_response(*[good_item() for _ in range(6)])
        with self.assertRaisesMessage(SuggestionError, "unexpected response format"):
            self.run_service()

    def test_code_fenced_json_is_accepted(self, urlopen):
        fenced = "```json\n" + json.dumps({"suggestions": [good_item()]}) + "\n```"
        urlopen.return_value = api_response(fenced)
        self.assertEqual(len(self.run_service()["suggestions"]), 1)

    def test_unreadable_response(self, urlopen):
        urlopen.return_value = api_response("not json at all")
        with self.assertRaisesMessage(SuggestionError, "unreadable"):
            self.run_service()

    def test_http_400_redacts_key_and_resume(self, urlopen):
        urlopen.side_effect = http_error(400, f"Problem with {FAKE_KEY} and text: {SOURCE}")
        with self.assertRaises(SuggestionError) as ctx:
            self.run_service()
        message = str(ctx.exception)
        self.assertIn("HTTP 400", message)
        self.assertNotIn(FAKE_KEY, message)
        self.assertNotIn("attendance tracker", message)

    def test_http_401_names_the_key(self, urlopen):
        urlopen.side_effect = http_error(401, "invalid x-api-key")
        with self.assertRaisesMessage(SuggestionError, "rejected the API key"):
            self.run_service()

    def test_disabled_makes_no_network_call(self, urlopen):
        with self.settings(CLAUDE_SUGGESTIONS_ENABLED=False):
            with self.assertRaisesMessage(SuggestionError, "disabled"):
                self.run_service()
        urlopen.assert_not_called()


@override_settings(CLAUDE_SUGGESTIONS_ENABLED=True, ANTHROPIC_API_KEY=FAKE_KEY)
class SuggestionViewTests(TestCase):
    def setUp(self):
        cache.clear()
        patches = {
            "extract": patch("analyser.api.extract_text", return_value="Projects\nBuilt a tracker using Python"),
            "score": patch("analyser.api.score_resume", return_value={
                "required_missing": ["docker"], "preferred_missing": [], "not_demonstrated": ["git"]}),
            "suggest": patch("analyser.api.create_suggestions", return_value={
                "suggestions": [], "draft_notice": "draft"}),
        }
        self.mocks = {name: p.start() for name, p in patches.items()}
        for p in patches.values():
            self.addCleanup(p.stop)

    def post(self, filename="resume.pdf", **extra):
        pdf = SimpleUploadedFile(filename, b"%PDF-1.4 fake", content_type="application/pdf")
        return self.client.post(SUGGEST_URL, {"resume": pdf, "job_description": "Python developer"}, **extra)

    def test_success_passes_skill_gaps_to_service(self):
        response = self.post()
        self.assertEqual(response.status_code, 200)
        self.assertIn("suggestions", response.json())
        self.mocks["suggest"].assert_called_once()
        self.assertEqual(self.mocks["suggest"].call_args[0][1], ["docker"])

    def test_provider_error_returns_502_with_message(self):
        self.mocks["suggest"].side_effect = SuggestionError("Anthropic returned HTTP 500.")
        response = self.post()
        self.assertEqual(response.status_code, 502)
        self.assertIn("HTTP 500", response.json()["error"])

    def test_disabled_returns_503_and_reads_nothing(self):
        with self.settings(CLAUDE_SUGGESTIONS_ENABLED=False):
            response = self.post()
        self.assertEqual(response.status_code, 503)
        self.mocks["extract"].assert_not_called()

    def test_disallowed_origin_returns_403(self):
        response = self.post(HTTP_ORIGIN="http://evil.example")
        self.assertEqual(response.status_code, 403)
        self.mocks["suggest"].assert_not_called()

    def test_rate_limit_after_five_requests(self):
        for _ in range(5):
            self.assertEqual(self.post().status_code, 200)
        self.assertEqual(self.post().status_code, 429)
        self.assertEqual(self.mocks["suggest"].call_count, 5)

    def test_non_pdf_returns_400(self):
        response = self.post(filename="notes.txt")
        self.assertEqual(response.status_code, 400)

    def test_invalid_upload_does_not_use_rate_limit(self):
        for _ in range(6):
            self.assertEqual(self.post(filename="notes.txt").status_code, 400)
        self.assertEqual(self.post().status_code, 200)    

class SectionParserTests(TestCase):
    def test_personal_details_are_separated(self):
        text = (
            "Certifications\n- Certification from Example Institute\n"
            "Extracurricular Activities\n- College events\n"
            "Personal Details\nFathers Name : Example Person\nAddress : Example Street"
        )
        sections = split_sections(text)
        self.assertIn("personal", sections)
        self.assertIn("Fathers Name", sections["personal"])
        self.assertNotIn("Fathers Name", sections["certifications"])
        self.assertNotIn("College events", sections["certifications"])    


class ApiBasicsTests(TestCase):
    def test_health(self):
        response = self.client.get("/api/health/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")

    def test_oversized_upload_rejected(self):
        big = SimpleUploadedFile("big.pdf", b"0" * (5 * 1024 * 1024 + 1), content_type="application/pdf")
        response = self.client.post("/api/analyse/", {"resume": big, "job_description": "Python"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("too large", response.json()["errors"]["resume"][0])


@patch("analyser.services.scorer.match_requirements", return_value=[])
class ScorerTests(TestCase):
    def test_matched_and_missing_skills(self, _):
        sections = {
            "summary": "Developer", "education": "BSc",
            "skills": "Python, SQL", "projects": "Built a tracker in Python",
        }
        result = score_resume(sections, "Required: Python, Django")
        self.assertEqual(result["required_matched"], ["python"])
        self.assertEqual(result["required_missing"], ["django"])

class DocxTests(TestCase):
    def test_docx_text_and_table_are_extracted(self):
        doc = WordDocument()
        doc.add_paragraph("Built a tracker using Python")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Skills: SQL, Git"
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        buffer.name = "resume.docx"

        text = extract_document_text(buffer)
        self.assertIn("Built a tracker using Python", text)
        self.assertIn("SQL, Git", text)        